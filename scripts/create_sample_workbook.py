#!/usr/bin/env python3
"""
create_sample_workbook.py
--------------------------
Generates input/sample_data_model.xlsx — a fully-populated example workbook
that demonstrates the exact sheet names, column headers, and data formats
expected by excel_to_docs.py.

The sample models a simple e-commerce domain:
  Customer → Order → OrderItem → Product, with a lookup and mapping sheet.

Usage:
    python scripts/create_sample_workbook.py
    python scripts/create_sample_workbook.py --output input/my_template.xlsx
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------------------
# Styling helpers
# ---------------------------------------------------------------------------

_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")   # dark blue
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
_SUBHEADER_FILL = PatternFill("solid", fgColor="2E75B6") # medium blue
_SUBHEADER_FONT = Font(color="FFFFFF", bold=True, size=9)
_EVEN_FILL = PatternFill("solid", fgColor="DEEAF1")
_ODD_FILL = PatternFill("solid", fgColor="FFFFFF")
_THIN_BORDER = Border(
    left=Side(style="thin", color="BDD7EE"),
    right=Side(style="thin", color="BDD7EE"),
    top=Side(style="thin", color="BDD7EE"),
    bottom=Side(style="thin", color="BDD7EE"),
)


def _style_header(ws, row: int, col_count: int) -> None:
    for c in range(1, col_count + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _THIN_BORDER


def _style_data_row(ws, row: int, col_count: int, even: bool) -> None:
    fill = _EVEN_FILL if even else _ODD_FILL
    for c in range(1, col_count + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = _THIN_BORDER


def _auto_width(ws) -> None:
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 4, 40)


def _write_sheet(ws, headers: list[str], rows: list[list]) -> None:
    ws.row_dimensions[1].height = 30
    ws.append(headers)
    _style_header(ws, 1, len(headers))
    ws.freeze_panes = "A2"

    for i, row in enumerate(rows):
        ws.append(row)
        _style_data_row(ws, i + 2, len(headers), i % 2 == 0)
        ws.row_dimensions[i + 2].height = 18

    _auto_width(ws)


# ---------------------------------------------------------------------------
# Sheet definitions
# ---------------------------------------------------------------------------

def _add_entities(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet("Entities")
    ws.sheet_properties.tabColor = "1F4E79"
    headers = ["Entity", "Description", "Domain", "Layer", "Notes"]
    rows = [
        ["Customer",  "A person or organisation that places orders.",          "Sales",    "mart",    ""],
        ["Order",     "A transaction placed by a customer.",                    "Sales",    "mart",    ""],
        ["OrderItem", "A single line in an order referencing a product.",       "Sales",    "mart",    ""],
        ["Product",   "A product available for sale.",                          "Catalogue","mart",    ""],
        ["OrderStatus","Lookup table for order status codes.",                  "Sales",    "staging", "Reference data"],
    ]
    _write_sheet(ws, headers, rows)


def _add_columns(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet("Columns")
    ws.sheet_properties.tabColor = "2E75B6"
    headers = [
        "Entity", "Column", "Type", "Nullable",
        "PK", "FK", "FK Reference",
        "Description", "Business Rules",
    ]
    rows = [
        # Customer
        ["Customer", "customer_id",    "INT64",     "No",  "Yes", "No",  "",                       "Surrogate key",               ""],
        ["Customer", "customer_code",  "STRING",    "No",  "No",  "No",  "",                       "Business key from source CRM","Must be unique"],
        ["Customer", "full_name",      "STRING",    "No",  "No",  "No",  "",                       "Full name",                   ""],
        ["Customer", "email",          "STRING",    "No",  "No",  "No",  "",                       "Primary email address",       "Must contain @"],
        ["Customer", "phone",          "STRING",    "Yes", "No",  "No",  "",                       "Contact phone number",        ""],
        ["Customer", "country_code",   "STRING",    "Yes", "No",  "No",  "",                       "ISO 3166-1 alpha-2",          "2 chars, uppercase"],
        ["Customer", "created_at",     "TIMESTAMP", "No",  "No",  "No",  "",                       "Record creation timestamp",   ""],
        ["Customer", "updated_at",     "TIMESTAMP", "Yes", "No",  "No",  "",                       "Last update timestamp",       ""],
        # Order
        ["Order",     "order_id",      "INT64",     "No",  "Yes", "No",  "",                       "Surrogate key",               ""],
        ["Order",     "order_ref",     "STRING",    "No",  "No",  "No",  "",                       "Business order reference",    "Format: ORD-YYYYNNNNNN"],
        ["Order",     "customer_id",   "INT64",     "No",  "No",  "Yes", "Customer.customer_id",   "FK to Customer",              ""],
        ["Order",     "status_code",   "STRING",    "No",  "No",  "Yes", "OrderStatus.code",       "FK to OrderStatus lookup",    ""],
        ["Order",     "order_date",    "DATE",      "No",  "No",  "No",  "",                       "Date order was placed",       ""],
        ["Order",     "total_amount",  "NUMERIC",   "No",  "No",  "No",  "",                       "Total order value (excl tax)","Must be >= 0"],
        ["Order",     "currency_code", "STRING",    "No",  "No",  "No",  "",                       "ISO 4217 currency code",      "3 chars, uppercase"],
        ["Order",     "created_at",    "TIMESTAMP", "No",  "No",  "No",  "",                       "Record creation timestamp",   ""],
        # OrderItem
        ["OrderItem", "order_item_id", "INT64",     "No",  "Yes", "No",  "",                       "Surrogate key",               ""],
        ["OrderItem", "order_id",      "INT64",     "No",  "No",  "Yes", "Order.order_id",         "FK to Order",                 ""],
        ["OrderItem", "product_id",    "INT64",     "No",  "No",  "Yes", "Product.product_id",     "FK to Product",               ""],
        ["OrderItem", "quantity",      "INT64",     "No",  "No",  "No",  "",                       "Number of units ordered",     "Must be > 0"],
        ["OrderItem", "unit_price",    "NUMERIC",   "No",  "No",  "No",  "",                       "Price per unit at time of order","Must be >= 0"],
        ["OrderItem", "line_total",    "NUMERIC",   "No",  "No",  "No",  "",                       "quantity * unit_price",       "Derived field"],
        # Product
        ["Product",   "product_id",    "INT64",     "No",  "Yes", "No",  "",                       "Surrogate key",               ""],
        ["Product",   "sku",           "STRING",    "No",  "No",  "No",  "",                       "Stock keeping unit",          "Unique, uppercase"],
        ["Product",   "product_name",  "STRING",    "No",  "No",  "No",  "",                       "Display name",                ""],
        ["Product",   "category",      "STRING",    "Yes", "No",  "No",  "",                       "Product category",            ""],
        ["Product",   "unit_cost",     "NUMERIC",   "Yes", "No",  "No",  "",                       "Cost price",                  ""],
        ["Product",   "is_active",     "BOOL",      "No",  "No",  "No",  "",                       "Whether product is on sale",  "Default TRUE"],
        ["Product",   "created_at",    "TIMESTAMP", "No",  "No",  "No",  "",                       "Record creation timestamp",   ""],
        # OrderStatus
        ["OrderStatus","code",         "STRING",    "No",  "Yes", "No",  "",                       "Status code",                 ""],
        ["OrderStatus","label",        "STRING",    "No",  "No",  "No",  "",                       "Human-readable label",        ""],
        ["OrderStatus","is_terminal",  "BOOL",      "No",  "No",  "No",  "",                       "No further transitions allowed",""],
    ]
    _write_sheet(ws, headers, rows)


def _add_relationships(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet("Relationships")
    ws.sheet_properties.tabColor = "375623"
    headers = ["From Entity", "To Entity", "Cardinality", "Label"]
    rows = [
        ["Customer",  "Order",     "1..N", "places"],
        ["Order",     "OrderItem", "1..N", "contains"],
        ["Product",   "OrderItem", "1..N", "included in"],
        ["OrderStatus","Order",    "1..N", "classifies"],
    ]
    _write_sheet(ws, headers, rows)


def _add_mappings(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet("Mappings")
    ws.sheet_properties.tabColor = "833C00"
    headers = [
        "Source System", "Source Table", "Source Column",
        "Target Entity", "Target Column",
        "Transformation", "Notes",
    ]
    rows = [
        ["CRM",   "crm_accounts", "account_id",    "Customer", "customer_code",  "CAST(account_id AS STRING)", "Natural key from CRM"],
        ["CRM",   "crm_accounts", "name",          "Customer", "full_name",      "TRIM(name)",                 ""],
        ["CRM",   "crm_accounts", "email_address", "Customer", "email",          "LOWER(email_address)",       "Normalise to lowercase"],
        ["CRM",   "crm_accounts", "phone_number",  "Customer", "phone",          "REGEXP_REPLACE(phone_number, '[^0-9+]', '')", "Strip non-numeric"],
        ["CRM",   "crm_accounts", "country",       "Customer", "country_code",   "UPPER(LEFT(country, 2))",    "ISO-2 from full name — verify mapping"],
        ["OMS",   "orders",       "id",            "Order",    "order_ref",      "CONCAT('ORD-', LPAD(CAST(id AS STRING), 10, '0'))", "Pad to 10 digits"],
        ["OMS",   "orders",       "cust_id",       "Order",    "customer_id",    "Lookup via customer_code",   "Join to Customer on code"],
        ["OMS",   "orders",       "status",        "Order",    "status_code",    "UPPER(status)",              ""],
        ["OMS",   "orders",       "placed_date",   "Order",    "order_date",     "DATE(placed_date)",          ""],
        ["OMS",   "order_lines",  "order_id",      "OrderItem","order_id",       "Direct",                     ""],
        ["OMS",   "order_lines",  "prod_sku",      "OrderItem","product_id",     "Lookup via Product.sku",     "Join to Product on sku"],
        ["OMS",   "order_lines",  "qty",           "OrderItem","quantity",       "CAST(qty AS INT64)",         ""],
        ["OMS",   "order_lines",  "price",         "OrderItem","unit_price",     "ROUND(price, 2)",            "2 decimal places"],
        ["PIM",   "products",     "sku",           "Product",  "sku",            "UPPER(TRIM(sku))",           ""],
        ["PIM",   "products",     "title",         "Product",  "product_name",   "TRIM(title)",                ""],
        ["PIM",   "products",     "category_name", "Product",  "category",       "Direct",                     ""],
        ["PIM",   "products",     "cost_price",    "Product",  "unit_cost",      "CAST(cost_price AS NUMERIC)",""],
        ["PIM",   "products",     "active_flag",   "Product",  "is_active",      "CAST(active_flag AS BOOL)",  "0/1 → FALSE/TRUE"],
    ]
    _write_sheet(ws, headers, rows)


def _add_lookups(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet("Lookups")
    ws.sheet_properties.tabColor = "7030A0"
    headers = ["Table Name", "Code", "Value", "Description"]
    rows = [
        ["OrderStatus", "PENDING",    "Pending",    "Order received, awaiting payment confirmation"],
        ["OrderStatus", "CONFIRMED",  "Confirmed",  "Payment confirmed, order being processed"],
        ["OrderStatus", "SHIPPED",    "Shipped",    "Order dispatched to carrier"],
        ["OrderStatus", "DELIVERED",  "Delivered",  "Order delivered to customer — terminal state"],
        ["OrderStatus", "CANCELLED",  "Cancelled",  "Order cancelled — terminal state"],
        ["OrderStatus", "REFUNDED",   "Refunded",   "Order refunded — terminal state"],
    ]
    _write_sheet(ws, headers, rows)


def _add_changelog(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet("ChangeLog")
    ws.sheet_properties.tabColor = "404040"
    headers = ["Version", "Date", "Author", "Description"]
    rows = [
        ["0.1", "2026-01-15", "J. Smith",   "Initial draft — Customer and Order entities"],
        ["0.2", "2026-02-03", "A. Lee",     "Added OrderItem and Product; first mapping pass"],
        ["0.3", "2026-03-01", "J. Smith",   "Added OrderStatus lookup; updated FK references"],
        ["1.0", "2026-03-18", "Data Team",  "First complete draft submitted for architecture review"],
    ]
    _write_sheet(ws, headers, rows)


def _add_notes(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet("Notes")
    ws.sheet_properties.tabColor = "FF0000"
    headers = ["#", "Type", "Description", "Owner", "Status"]
    rows = [
        ["1", "Open Question", "Should line_total be stored or always derived?",                    "Data Team",  "Open"],
        ["2", "Open Question", "Confirm ISO currency codes — CRM uses 3-char, OMS uses symbol",     "J. Smith",   "Open"],
        ["3", "TBD",           "Partitioning strategy for Order table — by order_date?",            "Platform",   "Open"],
        ["4", "TBD",           "dbt version to use — Core 1.8 or Cloud?",                          "Platform",   "Open"],
        ["5", "Assumption",    "customer_code is stable and can be used as a join key across systems","Data Team","Review"],
        ["6", "Note",          "PIM system does not provide created_at — will default to load time","A. Lee",     "Accepted"],
    ]
    _write_sheet(ws, headers, rows)
    # Add a legend row at the bottom
    ws.append([])
    ws.append(["", "Legend:", "Open = unresolved | TBD = needs decision | Assumption = flagged | Note = informational"])


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main() -> int:
    parser = argparse.ArgumentParser(
        description="Generate a sample Excel data model workbook for testing excel_to_docs.py."
    )
    parser.add_argument(
        "--output", default=None,
        help="Output path. Defaults to input/sample_data_model.xlsx relative to repo root.",
    )
    args = parser.parse_args()

    script_dir = Path(__file__).parent.resolve()
    repo_root = script_dir.parent

    if args.output:
        out_path = Path(args.output)
    else:
        out_path = repo_root / "input" / "sample_data_model.xlsx"

    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default empty sheet

    _add_entities(wb)
    _add_columns(wb)
    _add_relationships(wb)
    _add_mappings(wb)
    _add_lookups(wb)
    _add_changelog(wb)
    _add_notes(wb)

    wb.save(str(out_path))
    print(f"Sample workbook written to: {out_path}")
    print("\nSheets created:")
    for ws in wb.worksheets:
        print(f"  - {ws.title}")
    print("\nTo process this workbook:")
    print(f"  python scripts/excel_to_docs.py {out_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
